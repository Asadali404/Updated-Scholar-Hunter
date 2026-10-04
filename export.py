"""Excel export: Scholarships, Gap_Analysis, Profile, Summary."""
import io
from datetime import date, datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import tracker as tk

SCHOLARSHIP_COLS = ["id", "country", "scholarship_name", "provider", "level", "deadline", "requirements",
                    "funding", "official_link", "confidence", "fit_score", "status", "days_left",
                    "urgency", "notes"]


def records_df(records) -> pd.DataFrame:
    """Table used in the UI (Opportunities tab)."""
    return pd.DataFrame([{
        "Scholarship": r.name, "Provider": r.provider, "Country": r.country, "Level": r.level,
        "Funding": r.funding, "Deadline": r.deadline or "Not published", "Days left": r.days_left,
        "Urgency": tk.urgency(r.days_left), "Fit score": r.fit_score, "Confidence": r.confidence,
        "Source": "Web search" if r.source == "web" else "Curated list",
        "Verified": "Yes" if r.verified else "No", "Verification note": r.verification_note,
        "Official link": r.official_link, "Eligibility": r.eligibility,
    } for r in records])


def _conf(c: str) -> str:
    c = (c or "").lower()
    return c if c in ("high", "medium") else "low"      # Low / Estimated -> low


def _date(s: str):
    try:
        return datetime.strptime(s, "%Y-%m-%d").date() if s else None
    except ValueError:
        return None


def scholarships_df(records, tracker_df=None) -> pd.DataFrame:
    status, notes = {}, {}
    if tracker_df is not None and not tracker_df.empty:
        for _, row in tracker_df.iterrows():
            status[row["Scholarship"]] = row.get("Status", "Remaining")
            notes[row["Scholarship"]] = row.get("Notes", "") or ""
    rows = []
    for i, r in enumerate(records, start=1):
        rows.append({
            "id": i, "country": r.country, "scholarship_name": r.name, "provider": r.provider,
            "level": r.level, "deadline": _date(r.deadline), "requirements": r.requirements,
            "funding": r.funding, "official_link": r.official_link, "confidence": _conf(r.confidence),
            "fit_score": int(r.fit_score), "status": status.get(r.name, "Remaining"),
            "days_left": r.days_left, "urgency": tk.urgency_label(r.days_left),
            "notes": notes.get(r.name, ""),
        })
    return pd.DataFrame(rows, columns=SCHOLARSHIP_COLS)


def _style(ws, widths=None):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        longest = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[letter].width = min(max(12, longest + 2), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def to_excel(result, tracker_df=None) -> bytes:
    sch = scholarships_df(result.records, tracker_df)

    gap_rows = []
    if result.gap_report:
        gap_rows = [{"scholarship_name": i.scholarship, "weaknesses": "; ".join(i.weaknesses),
                     "recommendations": "; ".join(i.recommendations), "priority": i.priority,
                     "est_prep_time": i.prep_time} for i in result.gap_report.items]
    gap = pd.DataFrame(gap_rows, columns=["scholarship_name", "weaknesses", "recommendations",
                                          "priority", "est_prep_time"])

    p = result.profile
    prof_rows = []
    if p:
        prof_rows = [("degree", p.degree), ("cgpa", p.cgpa), ("publications", p.publications),
                     ("skills", ", ".join(p.skills)), ("research_areas", ", ".join(p.research_areas)),
                     ("keywords", ", ".join(p.keywords)), ("summary", p.summary)]
    prof = pd.DataFrame(prof_rows, columns=["key", "value"])

    counts = sch["status"].value_counts() if not sch.empty else {}
    critical = int(((sch["urgency"] == "Critical") & (sch["status"] != "Applied")).sum()) if not sch.empty else 0
    summary = pd.DataFrame([
        ("generated_on", date.today().isoformat()),
        ("mode", result.mode),
        ("total_scholarships", len(sch)),
        ("remaining", int(counts.get("Remaining", 0)) if len(sch) else 0),
        ("pending", int(counts.get("Pending", 0)) if len(sch) else 0),
        ("applied", int(counts.get("Applied", 0)) if len(sch) else 0),
        ("critical_deadlines", critical),
        ("verified_records", sum(1 for r in result.records if r.verified)),
        ("note", "Always verify deadlines and requirements on the official site."),
    ], columns=["metric", "value"])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        sch.to_excel(xw, sheet_name="Scholarships", index=False)
        gap.to_excel(xw, sheet_name="Gap_Analysis", index=False)
        prof.to_excel(xw, sheet_name="Profile", index=False)
        summary.to_excel(xw, sheet_name="Summary", index=False)
        for name in ("Scholarships", "Gap_Analysis", "Profile", "Summary"):
            _style(xw.sheets[name])

        ws = xw.sheets["Scholarships"]
        col = {c: i + 1 for i, c in enumerate(SCHOLARSHIP_COLS)}
        for row in range(2, ws.max_row + 1):
            link = ws.cell(row=row, column=col["official_link"])
            if link.value:
                link.hyperlink = str(link.value)
                link.font = Font(color="0563C1", underline="single")
            ws.cell(row=row, column=col["deadline"]).number_format = "yyyy-mm-dd"
        if ws.max_row > 1:
            dv = DataValidation(type="list", formula1='"Remaining,Pending,Applied"', allow_blank=False)
            ws.add_data_validation(dv)
            letter = ws.cell(row=1, column=col["status"]).column_letter
            dv.add(f"{letter}2:{letter}{ws.max_row}")
    return buf.getvalue()
